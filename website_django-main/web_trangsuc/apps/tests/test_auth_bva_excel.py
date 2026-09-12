"""
Excel-driven BVA tests cho Register / Login.

Doc: apps/tests/excel_data/auth_bva_tests.xlsx
Boundary rules: apps/auth_bounds.py

Chay:
  python manage.py test apps.tests.test_auth_bva_excel -v 2
"""
from __future__ import annotations

from pathlib import Path

from django.test import TestCase
from rest_framework.test import APIClient

from .excel_runner import (
    apply_setup,
    contains_expected,
    extract_response_data,
    load_excel_cases,
    resolve_case,
)

BVA_EXCEL_PATH = (
    Path(__file__).resolve().parent / "excel_data" / "auth_bva_tests.xlsx"
)


def _run_excel_case(test_case: TestCase, case: dict) -> None:
    client = APIClient()
    ctx = apply_setup(client, case["setup"])
    resolved = resolve_case(case, ctx)

    method = resolved["method"].lower()
    url = resolved["url"]
    payload = resolved["input"] or {}

    request_fn = getattr(client, method)
    if method in {"post", "put", "patch"}:
        response = request_fn(url, data=payload, format="json")
    else:
        response = request_fn(url)

    actual_status = response.status_code
    actual_body = extract_response_data(response)
    expected_status = resolved["expected_status"]
    expected_body = resolved["expected_result"]

    test_case.assertEqual(
        actual_status,
        expected_status,
        msg=(
            f"[{case['test_id']}] {case['description']}\n"
            f"boundary={case.get('boundary')} field={case.get('field')}\n"
            f"URL: {method.upper()} {url}\n"
            f"Input: {payload}\n"
            f"Expected status: {expected_status}, Actual: {actual_status}\n"
            f"Response: {actual_body}"
        ),
    )
    test_case.assertTrue(
        contains_expected(actual_body, expected_body),
        msg=(
            f"[{case['test_id']}] {case['description']}\n"
            f"Expected result (subset): {expected_body}\n"
            f"Actual response: {actual_body}"
        ),
    )


class AuthBvaExcelTest(TestCase):
    def test_bva_excel_file_exists(self):
        self.assertTrue(
            BVA_EXCEL_PATH.exists(),
            f"Thieu file Excel BVA: {BVA_EXCEL_PATH}",
        )
        cases = load_excel_cases(BVA_EXCEL_PATH)
        self.assertGreater(len(cases), 0)


def _attach_bva_tests() -> None:
    try:
        cases = load_excel_cases(BVA_EXCEL_PATH)
    except FileNotFoundError:
        return

    # giu them field/boundary tu Excel neu co
    from openpyxl import load_workbook

    wb = load_workbook(BVA_EXCEL_PATH, data_only=True)
    meta = {}
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in {"README", "BOUNDS"}:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            item = {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            tid = str(item.get("test_id", "")).strip()
            meta[tid] = {
                "field": str(item.get("field") or ""),
                "boundary": str(item.get("boundary") or ""),
            }

    used: set[str] = set()
    for case in cases:
        extra = meta.get(case["test_id"], {})
        case["field"] = extra.get("field", "")
        case["boundary"] = extra.get("boundary", "")

        base = f"test_{case['sheet']}_{case['test_id']}".replace("-", "_")
        name = base
        i = 2
        while name in used or hasattr(AuthBvaExcelTest, name):
            name = f"{base}_{i}"
            i += 1
        used.add(name)

        def _make(c):
            def _test(self):
                _run_excel_case(self, c)

            _test.__doc__ = f"{c['test_id']}: {c['description']}"
            _test.__name__ = name
            return _test

        setattr(AuthBvaExcelTest, name, _make(case))


_attach_bva_tests()
