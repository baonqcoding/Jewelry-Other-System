"""
Excel-driven BVA + Equivalence Partitioning cho toan bo backend (khong UI).

Doc: apps/tests/excel_data/backend_bva_ep_tests.xlsx
Bounds: apps/auth_bounds.py, apps/domain_bounds.py

Chay:
  python manage.py test apps.tests.test_bva_ep_excel -v 2
"""
from __future__ import annotations

from pathlib import Path

from django.test import TestCase
from rest_framework.test import APIClient

from .excel_runner import (
    apply_setup,
    contains_expected,
    extract_response_data,
    load_excel_cases,
    resolve_case,
)

BVA_EP_EXCEL = (
    Path(__file__).resolve().parent / "excel_data" / "backend_bva_ep_tests.xlsx"
)


def _run_case(test_case: TestCase, case: dict) -> None:
    client = APIClient()
    ctx = apply_setup(client, case["setup"])
    resolved = resolve_case(case, ctx)

    method = resolved["method"].lower()
    url = resolved["url"]
    payload = resolved["input"] or {}

    request_fn = getattr(client, method)
    if method in {"post", "put", "patch"}:
        response = request_fn(url, data=payload, format="json")
    else:
        response = request_fn(url)

    actual_status = response.status_code
    actual_body = extract_response_data(response)

    test_case.assertEqual(
        actual_status,
        resolved["expected_status"],
        msg=(
            f"[{case['test_id']}] {case['description']}\n"
            f"{case.get('technique')} | {case.get('field')} | {case.get('partition_or_boundary')}\n"
            f"{method.upper()} {url}\nInput: {payload}\n"
            f"Expected {resolved['expected_status']}, got {actual_status}\n"
            f"Body: {actual_body}"
        ),
    )
    test_case.assertTrue(
        contains_expected(actual_body, resolved["expected_result"]),
        msg=(
            f"[{case['test_id']}] expected subset {resolved['expected_result']}\n"
            f"actual: {actual_body}"
        ),
    )


class BackendBvaEpExcelTest(TestCase):
    def test_excel_loaded(self):
        self.assertTrue(BVA_EP_EXCEL.exists(), f"Missing {BVA_EP_EXCEL}")
        self.assertGreater(len(load_excel_cases(BVA_EP_EXCEL)), 0)


def _attach() -> None:
    try:
        cases = load_excel_cases(BVA_EP_EXCEL)
    except FileNotFoundError:
        return

    # bo sung meta technique/field tu Excel
    from openpyxl import load_workbook

    wb = load_workbook(BVA_EP_EXCEL, data_only=True)
    meta = {}
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in {"README", "BOUNDS", "ALL_CASES"}:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            item = {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            tid = str(item.get("test_id") or "").strip()
            if not tid:
                continue
            meta[tid] = {
                "technique": str(item.get("technique") or ""),
                "field": str(item.get("field") or ""),
                "partition_or_boundary": str(item.get("partition_or_boundary") or ""),
            }

    used: set[str] = set()
    for case in cases:
        case.update(meta.get(case["test_id"], {}))
        base = f"test_{case['sheet']}_{case['test_id']}".replace("-", "_")
        name = base
        i = 2
        while name in used or hasattr(BackendBvaEpExcelTest, name):
            name = f"{base}_{i}"
            i += 1
        used.add(name)

        def _make(c):
            def _test(self):
                _run_case(self, c)

            _test.__doc__ = f"{c['test_id']}: {c['description']}"
            _test.__name__ = name
            return _test

        setattr(BackendBvaEpExcelTest, name, _make(case))


_attach()
